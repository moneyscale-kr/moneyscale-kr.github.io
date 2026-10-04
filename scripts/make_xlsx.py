#!/usr/bin/env python3
"""월급 관리 엑셀 lite 생성. 사용: python3 scripts/make_xlsx.py (openpyxl 필요)"""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

root = pathlib.Path(__file__).resolve().parent.parent
INK, LIME, CORAL, PAPER, GRID = "0B0B0F", "C8FF3D", "FF5A4E", "F4F1EA", "D9D6CC"
F = "Malgun Gothic"
title_f = Font(name=F, size=16, bold=True, color=LIME)
head_f = Font(name=F, size=11, bold=True, color=INK)
body_f = Font(name=F, size=11, color="222222")
in_f = Font(name=F, size=12, bold=True, color="0B5FFF")      # 입력칸 = 파란 글씨
note_f = Font(name=F, size=9, color="777777")
warn_f = Font(name=F, size=9, color="C0392B")
dark = PatternFill("solid", fgColor=INK)
lime = PatternFill("solid", fgColor=LIME)
inp = PatternFill("solid", fgColor="FFF9DB")
soft = PatternFill("solid", fgColor="F4F1EA")
thin = Side(style="thin", color=GRID); box = Border(left=thin, right=thin, top=thin, bottom=thin)
WON = '#,##0"원"'; PCT = "0.0%"

def banner(ws, text, cols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    c = ws.cell(1, 1, text); c.font = title_f; c.fill = dark; c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    ws.sheet_view.showGridLines = False
def hdr(ws, r, vals, c0=1):
    for i, v in enumerate(vals):
        c = ws.cell(r, c0 + i, v); c.font = head_f; c.fill = lime; c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center")
def cell(ws, r, c, v, fmt=None, font=body_f, fill=None, align=None):
    x = ws.cell(r, c, v); x.font = font; x.border = box
    if fmt: x.number_format = fmt
    if fill: x.fill = fill
    if align: x.alignment = Alignment(horizontal=align, vertical="center")
    return x

wb = Workbook()
# ---------------- 시트 1: 월급 배분 ----------------
s1 = wb.active; s1.title = "1_월급배분"
banner(s1, "돈의 스케일 | 월급 관리 엑셀 lite  -  1. 월급 배분", 5)
for col, w in zip("ABCDE", (22, 18, 14, 18, 40)): s1.column_dimensions[col].width = w
s1["A3"] = "노란 칸(파란 글씨)만 바꾸세요. 나머지는 자동 계산돼요."; s1["A3"].font = note_f
cell(s1, 5, 1, "월 실수령액", font=head_f, fill=soft); cell(s1, 5, 2, 2800000, WON, in_f, inp)
s1["C5"] = "통장에 실제로 찍히는 금액"; s1["C5"].font = note_f
hdr(s1, 7, ["구분", "내 목표 비율", "금액", "6개월 누적", "메모"])
rows = [("고정비 (월세·관리비·통신·보험료 등)", 0.50, "매달 같은 날 나가는 돈"),
        ("생활비 (식비·교통·여가)", 0.20, "변동비. 한도를 정해두면 편해요"),
        ("저축 (비상금·예금·적금)", 0.20, "원금이 보장되는 쪽 위주로 정해보세요"),
        ("투자·자기계발 (본인 계획)", 0.10, "비율은 예시일 뿐 권유가 아니에요")]
for i, (n, p, m) in enumerate(rows):
    r = 8 + i
    cell(s1, r, 1, n); cell(s1, r, 2, p, "0%", in_f, inp, "center")
    cell(s1, r, 3, f"=ROUND($B$5*B{r},-1)", WON); cell(s1, r, 4, f"=C{r}*6", WON); cell(s1, r, 5, m, font=note_f)
cell(s1, 12, 1, "합계", font=head_f, fill=soft); cell(s1, 12, 2, "=ROUND(SUM(B8:B11),4)", "0%", head_f, soft, "center")
cell(s1, 12, 3, "=SUM(C8:C11)", WON, head_f, soft); cell(s1, 12, 4, "=SUM(D8:D11)", WON, head_f, soft)
cell(s1, 12, 5, '=IF(ROUND(B12,4)=1,"OK: 비율 합계 100%","비율 합계가 100%가 아니에요")', font=warn_f)
s1.conditional_formatting.add("B12", CellIsRule(operator="notEqual", formula=["1"], font=Font(color=CORAL, bold=True)))
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", showErrorMessage=True,
                    errorTitle="비율", error="0%~100% 사이로 입력하세요"); s1.add_data_validation(dv); dv.add("B8:B11")
s1["A14"] = "저축 가능 빠른 계산"; s1["A14"].font = Font(name=F, size=12, bold=True)
cell(s1, 15, 1, "연 저축액"); cell(s1, 15, 2, "=C10*12", WON)
cell(s1, 16, 1, "1억 원까지 걸리는 기간(년)"); cell(s1, 16, 2, '=IF(C10>0,ROUND(100000000/(C10*12),1),"-")', "0.0")
s1["C16"] = "이자·세금 미반영 단순 계산"; s1["C16"].font = note_f
s1["A18"] = "참고용 템플릿이에요. 세법·보험요율은 변경될 수 있고, 특정 금융상품·투자에 대한 권유가 아닙니다."; s1["A18"].font = warn_f
s1["A19"] = "비율 예시는 일반적인 가계 관리 틀이며, 개인 상황에 맞게 바꿔 쓰세요.  moneyscale-kr.github.io"; s1["A19"].font = note_f

# ---------------- 시트 2: 고정비 목록 ----------------
s2 = wb.create_sheet("2_고정비_저축목록")
banner(s2, "2. 고정비 · 저축 목록  -  항목을 직접 적으세요", 5)
for col, w in zip("ABCDE", (26, 16, 14, 16, 30)): s2.column_dimensions[col].width = w
hdr(s2, 3, ["항목", "월 금액", "분류", "나가는 날(일)", "메모"])
items = [("월세/대출 상환(본인 계약)", 700000, "고정비", 1), ("관리비·공과금", 120000, "고정비", 10),
         ("통신비", 55000, "고정비", 15), ("교통(정기권)", 60000, "고정비", 20), ("구독 서비스", 30000, "고정비", 25),
         ("비상금 적립", 300000, "저축", 25), ("정기 저축", 300000, "저축", 25), ("자기계발", 100000, "투자·자기계발", 28)]
for i in range(20):
    r = 4 + i
    for c in range(1, 6): cell(s2, r, c, None, fill=inp if c in (1, 2, 3, 4) else None)
    if i < len(items):
        n, a, k, d = items[i]; s2.cell(r, 1, n); s2.cell(r, 2, a); s2.cell(r, 3, k); s2.cell(r, 4, d)
    s2.cell(r, 2).number_format = WON; s2.cell(r, 2).font = in_f; s2.cell(r, 1).font = in_f
    s2.cell(r, 3).alignment = Alignment(horizontal="center"); s2.cell(r, 4).alignment = Alignment(horizontal="center")
dv2 = DataValidation(type="list", formula1='"고정비,생활비,저축,투자·자기계발"', allow_blank=True); s2.add_data_validation(dv2); dv2.add("C4:C23")
hdr(s2, 25, ["분류별 합계", "목록 합계", "목표 금액", "차이", "판정"])
for i, k in enumerate(["고정비", "생활비", "저축", "투자·자기계발"]):
    r = 26 + i
    cell(s2, r, 1, k, font=head_f, fill=soft)
    cell(s2, r, 2, f'=SUMIF($C$4:$C$23,A{r},$B$4:$B$23)', WON)
    cell(s2, r, 3, f"='1_월급배분'!C{8+i}", WON)
    cell(s2, r, 4, f"=C{r}-B{r}", '#,##0"원";[Red]-#,##0"원"')
    cell(s2, r, 5, f'=IF(B{r}=0,"목록 비어 있음",IF(B{r}<=C{r},"목표 안","목표 초과"))', align="center")
s2.conditional_formatting.add("E26:E29", CellIsRule(operator="equal", formula=['"목표 초과"'], font=Font(color=CORAL, bold=True)))
s2["A31"] = "고정비 비중 = 고정비 합계 / 월 실수령액"; s2["A31"].font = Font(name=F, size=11, bold=True)
cell(s2, 31, 2, "=IFERROR(B26/'1_월급배분'!B5,0)", PCT, head_f, soft)
s2["A33"] = "고정비가 실수령의 절반을 넘으면 항목을 점검해 보세요. 참고용이며 개인 상황에 따라 다릅니다."; s2["A33"].font = note_f

# ---------------- 시트 3: 12개월 기록 ----------------
s3 = wb.create_sheet("3_12개월기록")
banner(s3, "3. 12개월 기록  -  매달 실제 금액을 적고 비율을 확인하세요", 9)
for col, w in zip("ABCDEFGHI", (10, 16, 14, 14, 14, 14, 14, 12, 12)): s3.column_dimensions[col].width = w
hdr(s3, 3, ["월", "실수령", "고정비", "생활비", "저축", "투자·자기계발", "남은 돈", "고정비%", "저축%"])
for m in range(1, 13):
    r = 3 + m
    cell(s3, r, 1, f"{m}월", font=head_f, fill=soft, align="center")
    for c in range(2, 7): cell(s3, r, c, None, WON, in_f, inp)
    cell(s3, r, 7, f'=IF(B{r}="","",B{r}-SUM(C{r}:F{r}))', '#,##0"원";[Red]-#,##0"원"')
    cell(s3, r, 8, f'=IF(N(B{r})=0,"",C{r}/B{r})', PCT, align="center")
    cell(s3, r, 9, f'=IF(N(B{r})=0,"",E{r}/B{r})', PCT, align="center")
cell(s3, 16, 1, "합계", font=head_f, fill=soft, align="center")
for c, L in zip(range(2, 8), "BCDEFG"): cell(s3, 16, c, f"=SUM({L}4:{L}15)", WON, head_f, soft)
cell(s3, 16, 8, '=IF(B16=0,"",C16/B16)', PCT, head_f, soft, "center"); cell(s3, 16, 9, '=IF(B16=0,"",E16/B16)', PCT, head_f, soft, "center")
s3["A18"] = "남은 돈이 마이너스(빨강)면 그 달 지출이 실수령보다 컸다는 뜻이에요."; s3["A18"].font = note_f
s3["A19"] = "참고용 템플릿이에요. 투자 권유·금융상품 추천이 아닙니다.  돈의 스케일 moneyscale-kr.github.io"; s3["A19"].font = warn_f
s3.freeze_panes = "B4"

for ws in (s1, s2, s3):
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
    ws.sheet_properties.tabColor = LIME
wb.properties.title = "월급 관리 엑셀 lite"; wb.properties.creator = "돈의 스케일"
out = root / "downloads" / "money-scale-salary-lite.xlsx"; wb.save(out); print("saved", out)
